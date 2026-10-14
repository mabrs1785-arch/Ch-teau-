#!/usr/bin/env python3
"""Générateur du modèle financier 10 ans — Opération Château.

Chaque feuille de scénario contient un bloc d'hypothèses éditables (colonne B, en bleu)
et un calcul sur 10 ans entièrement en formules Excel : modifier une hypothèse dans le
fichier .xlsx recalcule tout. Ce script ne sert qu'à (re)générer le fichier.

Usage : python3 build_model.py  (lit hypotheses.py, écrit Previsionnels_10_ans.xlsx)
"""
import copy
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import hypotheses as H  # noqa: E402

N = 10
YC = [get_column_letter(3 + i) for i in range(N)]  # colonnes C..L = années 1..10
BLUE = Font(color="0000CC")
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="DDE7F0")
INP = PatternFill("solid", fgColor="FFF6D5")
EUR = '#,##0;[Red]-#,##0'
PCT = '0.0%'


class Sheet:
    def __init__(self, ws):
        self.ws = ws
        self.row = 1
        self.ref = {}

    def put(self, r, c, v, fmt=None, font=None, fill=None):
        cell = self.ws.cell(row=r, column=c, value=v)
        if fmt:
            cell.number_format = fmt
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        return cell

    def title(self, text):
        self.put(self.row, 1, text, font=Font(bold=True, size=13))
        self.row += 1

    def section(self, text):
        self.row += 1
        for c in range(1, 4 + N):
            self.ws.cell(row=self.row, column=c).fill = HEAD
        self.put(self.row, 1, text, font=BOLD, fill=HEAD)
        self.row += 1

    def inp(self, key, label, value, unit="", source="", fmt=None):
        r = self.row
        self.put(r, 1, label)
        if fmt is None:
            fmt = PCT if unit == "%" else ('#,##0.00' if isinstance(value, float) and abs(value) < 10 else EUR)
        self.put(r, 2, value, fmt=fmt, font=BLUE, fill=INP)
        self.put(r, 3, unit)
        self.put(r, 4, source)
        self.ref[key] = f"$B${r}"
        self.row += 1

    def years_header(self):
        r = self.row
        self.put(r, 1, "Année", font=BOLD)
        self.put(r, 2, "Total / réf.", font=BOLD)
        for i, c in enumerate(YC):
            self.put(r, 3 + i, i + 1, font=BOLD)
        self.ref["yr"] = r
        self.row += 1

    def line(self, key, label, fn, fmt=EUR, total=True, bold=False):
        """fn(i, col) -> formule (sans '=') pour l'année i (1..10) en colonne col."""
        r = self.row
        self.put(r, 1, label, font=BOLD if bold else None)
        for i, c in enumerate(YC):
            self.put(r, 3 + i, "=" + fn(i + 1, c), fmt=fmt, font=BOLD if bold else None)
        if total:
            self.put(r, 2, f"=SUM(C{r}:L{r})", fmt=fmt, font=BOLD if bold else None)
        self.ref[key] = r
        self.row += 1
        return r

    def R(self, key, col):
        return f"{col}{self.ref[key]}"

    def prev(self, col):
        idx = YC.index(col)
        return YC[idx - 1] if idx > 0 else None


def build_scenario(wb, name, p):
    ws = wb.create_sheet(name)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 15
    for c in YC:
        ws.column_dimensions[c].width = 12
    ws.column_dimensions["D"].width = 14
    s = Sheet(ws)
    s.title(p["titre"])
    s.put(s.row, 1, "Cellules jaunes/bleues = hypothèses éditables. Tout le reste est calculé. "
          "Colonne D (bloc hypothèses) = source ou statut (ESTIMATION = non vérifiée).",
          font=Font(italic=True))
    s.row += 1
    yr_row_placeholder = None  # noqa: F841

    s.section("1. Calendrier, acquisition et travaux")
    for k, lab, unit, src in [
        ("annee_acq", "Année d'acquisition des murs (1 = immédiate)", "année", ""),
        ("annee_ouv", "Année d'ouverture commerciale (avant stress)", "année", ""),
        ("prix", "Prix d'acquisition", "€", ""),
        ("frais_acq_pct", "Frais d'acquisition (droits + notaire)", "%", ""),
        ("part_bati", "Part du prix affectable au bâti (amortissable)", "%", ""),
        ("capex_restauration", "Travaux de restauration / mise aux normes", "€", ""),
        ("capex_equestre", "Investissements équestres (boxes, carrière, manège…)", "€", ""),
        ("capex_evenementiel", "Investissements événementiel / hébergement", "€", ""),
        ("capex_rep1", "Répartition travaux : année d'acquisition", "%", ""),
        ("capex_rep2", "Répartition travaux : année +1", "%", ""),
        ("capex_rep3", "Répartition travaux : année +2", "%", ""),
        ("gros_entretien_pct", "Gros entretien patrimonial annuel (% prix+travaux)", "%", ""),
        ("valo_sortie_pct", "Valeur de sortie an 10 (% prix + travaux cumulés)", "%", ""),
        ("loyer_avant_acq", "Loyer annuel versé avant acquisition (scénario location)", "€", ""),
    ]:
        v, sc = p["inputs"][k]
        s.inp(k, lab, v, unit, sc)

    s.section("2. Financement")
    for k, lab, unit in [
        ("apport_fondateur", "Apport du fondateur (année 1)", "€"),
        ("fp_invest", "Fonds propres investisseurs (année d'acquisition)", "€"),
        ("subventions", "Subventions (non acquises : 0 par défaut)", "€"),
        ("dette", "Dette bancaire (année d'acquisition)", "€"),
        ("taux", "Taux de la dette bancaire", "%"),
        ("duree", "Durée dette bancaire", "ans"),
        ("differe", "Différé d'amortissement (dans la durée)", "ans"),
        ("cv", "Crédit-vendeur", "€"),
        ("cv_taux", "Taux crédit-vendeur", "%"),
        ("cv_duree", "Durée crédit-vendeur (amortissable)", "ans"),
        ("oc", "Obligations / quasi-fonds propres (in fine)", "€"),
        ("oc_taux", "Coupon annuel des obligations", "%"),
        ("oc_ech", "Année de remboursement des obligations", "année"),
        ("part_fondateur", "Part du capital détenue par le fondateur", "%"),
        ("votes_fondateur", "Droits de vote du fondateur", "%"),
        ("treso_min", "Réserve de liquidité cible", "€"),
    ]:
        v, sc = p["inputs"][k]
        s.inp(k, lab, v, unit, sc)

    s.section("3. Paramètres de stress (neutres = 100 % / 0)")
    for k, lab, unit in [
        ("ca_mult", "Multiplicateur de chiffre d'affaires", "%"),
        ("capex_mult", "Multiplicateur du budget travaux", "%"),
        ("delai", "Retard d'ouverture", "ans"),
        ("taux_add", "Hausse du taux bancaire (points)", "%"),
        ("retrait_inv", "Part des fonds investisseurs non versée", "%"),
        ("annul_evt", "Part des événements annulés", "%"),
    ]:
        v = p["stress"][k]
        s.inp(k, lab, v, unit, "Stress test" if v not in (1, 0) else "neutre")

    s.section("4. Exploitation — hypothèses générales")
    for k, lab, unit in [
        ("ramp1", "Montée en charge : 1re année d'exploitation", "%"),
        ("ramp2", "Montée en charge : 2e année", "%"),
        ("ramp3", "Montée en charge : 3e année", "%"),
        ("infl_prix", "Indexation annuelle des prix", "%"),
        ("infl_ch", "Inflation annuelle des charges", "%"),
        ("bfr_pct", "BFR en % du CA (négatif = acomptes clients)", "%"),
        ("is_taux", "Taux d'IS (simplifié, taux normal)", "%"),
    ]:
        v, sc = p["inputs"][k]
        s.inp(k, lab, v, unit, sc)

    s.section("5. Lignes de revenus (à maturité, euros de l'année d'ouverture)")
    s.put(s.row, 1, "Libellé", font=BOLD)
    s.put(s.row, 2, "Volume", font=BOLD)
    s.put(s.row, 3, "Prix unitaire", font=BOLD)
    s.put(s.row, 4, "Coût var. %", font=BOLD)
    s.put(s.row, 5, "Événement ?", font=BOLD)
    s.put(s.row, 6, "Équestre ?", font=BOLD)
    s.put(s.row, 7, "Unité / source", font=BOLD)
    s.row += 1
    rev_rows = []
    for (lab, vol, prix, cv, evt, equ, src) in p["revenus"]:
        r = s.row
        s.put(r, 1, lab)
        s.put(r, 2, vol, fmt='#,##0', font=BLUE, fill=INP)
        s.put(r, 3, prix, fmt=EUR, font=BLUE, fill=INP)
        s.put(r, 4, cv, fmt=PCT, font=BLUE, fill=INP)
        s.put(r, 5, 1 if evt else 0, font=BLUE, fill=INP)
        s.put(r, 6, 1 if equ else 0, font=BLUE, fill=INP)
        s.put(r, 7, src)
        rev_rows.append(r)
        s.row += 1

    s.section("6. Charges fixes annuelles (euros de l'année d'ouverture)")
    s.put(s.row, 1, "Libellé", font=BOLD)
    s.put(s.row, 2, "Montant/an", font=BOLD)
    s.put(s.row, 3, "Dès l'acq. (1) / dès l'ouverture (0)", font=BOLD)
    s.put(s.row, 7, "Source / statut", font=BOLD)
    s.row += 1
    fix_rows = []
    for (lab, amt, from_acq, src) in p["charges"]:
        r = s.row
        s.put(r, 1, lab)
        s.put(r, 2, amt, fmt=EUR, font=BLUE, fill=INP)
        s.put(r, 3, 1 if from_acq else 0, font=BLUE, fill=INP)
        s.put(r, 7, src)
        fix_rows.append(r)
        s.row += 1

    I = s.ref  # noqa: E741
    s.section("7. Calcul sur 10 ans")
    s.years_header()
    yr = I["yr"]

    def y(col):
        return f"{col}${yr}"

    # Calendrier
    s.line("ouv_eff", "Année d'ouverture effective", lambda i, c: f"{I['annee_ouv']}+{I['delai']}", fmt='0', total=False)
    s.line("n_exp", "Rang de l'année d'exploitation", lambda i, c: f"MAX(0,{y(c)}-{s.R('ouv_eff', c)}+1)", fmt='0', total=False)
    s.line("ramp", "Coefficient de montée en charge",
           lambda i, c: f"IF({s.R('n_exp', c)}=0,0,IF({s.R('n_exp', c)}=1,{I['ramp1']},IF({s.R('n_exp', c)}=2,{I['ramp2']},IF({s.R('n_exp', c)}=3,{I['ramp3']},1))))",
           fmt=PCT, total=False)
    s.line("idx_p", "Indice prix", lambda i, c: f"(1+{I['infl_prix']})^({y(c)}-1)", fmt='0.000', total=False)
    s.line("idx_c", "Indice charges", lambda i, c: f"(1+{I['infl_ch']})^({y(c)}-1)", fmt='0.000', total=False)
    s.line("possede", "Murs détenus (1/0)", lambda i, c: f"IF({y(c)}>={I['annee_acq']},1,0)", fmt='0', total=False)

    # Revenus
    s.row += 1
    s.put(s.row, 1, "Compte de résultat", font=BOLD)
    s.row += 1
    rev_line_rows = []
    for k, r in enumerate(rev_rows):
        lab = ws.cell(row=r, column=1).value
        rr = s.line(f"rev{k}", "  " + lab,
                    lambda i, c, r=r: f"$B${r}*$C${r}*{s.R('ramp', c)}*{s.R('idx_p', c)}*{I['ca_mult']}*(1-$E${r}*{I['annul_evt']})")
        rev_line_rows.append((rr, r))
    s.line("ca", "Chiffre d'affaires total",
           lambda i, c: "+".join(f"{c}{rr}" for rr, _ in rev_line_rows), bold=True)
    s.line("ca_equ", "  dont équestre",
           lambda i, c: "+".join(f"{c}{rr}*$F${r}" for rr, r in rev_line_rows))
    s.line("part_equ", "  Part équestre du CA",
           lambda i, c: f"IF({s.R('ca', c)}=0,0,{s.R('ca_equ', c)}/{s.R('ca', c)})", fmt=PCT, total=False)
    s.line("cvar", "Coûts variables",
           lambda i, c: "-(" + "+".join(f"{c}{rr}*$D${r}" for rr, r in rev_line_rows) + ")")
    s.line("marge", "Marge sur coûts variables", lambda i, c: f"{s.R('ca', c)}+{s.R('cvar', c)}", bold=True)
    fix_line_rows = []
    for k, r in enumerate(fix_rows):
        lab = ws.cell(row=r, column=1).value
        fr = s.line(f"fix{k}", "  " + lab,
                    lambda i, c, r=r: f"-$B${r}*{s.R('idx_c', c)}*IF($C${r}=1,{s.R('possede', c)},IF({s.R('n_exp', c)}>0,1,0))")
        fix_line_rows.append(fr)
    s.line("fixes", "Charges fixes totales", lambda i, c: "+".join(f"{c}{fr}" for fr in fix_line_rows), bold=True)
    s.line("loyer", "Loyer des murs (avant acquisition)",
           lambda i, c: f"-IF({s.R('possede', c)}=0,IF({s.R('n_exp', c)}>0,{I['loyer_avant_acq']}*{s.R('idx_c', c)},0),0)")
    s.line("ebitda", "EBITDA", lambda i, c: f"{s.R('marge', c)}+{s.R('fixes', c)}+{s.R('loyer', c)}", bold=True)
    s.line("marge_ebitda", "  Marge d'EBITDA", lambda i, c: f"IF({s.R('ca', c)}=0,0,{s.R('ebitda', c)}/{s.R('ca', c)})", fmt=PCT, total=False)

    # Investissements (pour amortissements)
    s.line("capex_init", "Travaux et équipements initiaux",
           lambda i, c: f"({I['capex_restauration']}+{I['capex_equestre']}+{I['capex_evenementiel']})*{I['capex_mult']}*"
                        f"IF({y(c)}={I['annee_acq']},{I['capex_rep1']},IF({y(c)}={I['annee_acq']}+1,{I['capex_rep2']},IF({y(c)}={I['annee_acq']}+2,{I['capex_rep3']},0)))*{s.R('idx_c', c)}")
    s.line("acq", "Acquisition (prix + frais)",
           lambda i, c: f"IF({y(c)}={I['annee_acq']},{I['prix']}*(1+{I['frais_acq_pct']}),0)")
    s.line("ge", "Gros entretien patrimonial récurrent",
           lambda i, c: f"IF({y(c)}>={I['annee_acq']}+3,{I['gros_entretien_pct']}*({I['prix']}+({I['capex_restauration']}+{I['capex_equestre']}+{I['capex_evenementiel']})*{I['capex_mult']})*{s.R('idx_c', c)},0)")
    s.line("cum_capex", "Investissements amortissables cumulés",
           lambda i, c: (f"{s.R('capex_init', c)}+{s.R('ge', c)}" if i == 1 else
                         f"{s.R('cum_capex', s.prev(c))}+{s.R('capex_init', c)}+{s.R('ge', c)}"), total=False)
    s.line("dota", "Dotations aux amortissements (bâti 30 ans, travaux 15 ans)",
           lambda i, c: f"-({s.R('possede', c)}*{I['prix']}*(1+{I['frais_acq_pct']})*{I['part_bati']}/30+{s.R('cum_capex', c)}/15)")
    s.line("ebit", "Résultat d'exploitation", lambda i, c: f"{s.R('ebitda', c)}+{s.R('dota', c)}", bold=True)

    # Dette bancaire
    s.row += 1
    s.put(s.row, 1, "Échéancier des emprunts", font=BOLD)
    s.row += 1
    s.line("taux_eff", "Taux bancaire effectif", lambda i, c: f"{I['taux']}+{I['taux_add']}", fmt=PCT, total=False)
    s.line("k", "Rang de l'année d'emprunt", lambda i, c: f"MAX(0,{y(c)}-{I['annee_acq']}+1)", fmt='0', total=False)
    s.line("d_deb", "Dette bancaire — capital début",
           lambda i, c: (f"IF({s.R('k', c)}=1,{I['dette']},0)" if i == 1 else
                         f"IF({s.R('k', c)}=1,{I['dette']},{s.R('d_fin', s.prev(c))})"), total=False)
    s.line("d_int", "  Intérêts", lambda i, c: f"{s.R('d_deb', c)}*{s.R('taux_eff', c)}")
    s.line("d_amo", "  Amortissement du capital",
           lambda i, c: f"IF(OR({s.R('k', c)}=0,{s.R('k', c)}<={I['differe']},{s.R('d_deb', c)}<=0),0,"
                        f"MIN({s.R('d_deb', c)},PMT({s.R('taux_eff', c)},{I['duree']}-{I['differe']},-{I['dette']})-{s.R('d_int', c)}))")
    s.line("d_fin", "Dette bancaire — capital fin", lambda i, c: f"{s.R('d_deb', c)}-{s.R('d_amo', c)}", total=False)
    s.line("cv_deb", "Crédit-vendeur — capital début",
           lambda i, c: (f"IF({s.R('k', c)}=1,{I['cv']},0)" if i == 1 else
                         f"IF({s.R('k', c)}=1,{I['cv']},{s.R('cv_fin', s.prev(c))})"), total=False)
    s.line("cv_int", "  Intérêts crédit-vendeur", lambda i, c: f"{s.R('cv_deb', c)}*{I['cv_taux']}")
    s.line("cv_amo", "  Amortissement crédit-vendeur",
           lambda i, c: f"IF(OR({s.R('k', c)}=0,{s.R('cv_deb', c)}<=0),0,MIN({s.R('cv_deb', c)},IF({I['cv_taux']}=0,{I['cv']}/{I['cv_duree']},PMT({I['cv_taux']},{I['cv_duree']},-{I['cv']})-{s.R('cv_int', c)})))")
    s.line("cv_fin", "Crédit-vendeur — capital fin", lambda i, c: f"{s.R('cv_deb', c)}-{s.R('cv_amo', c)}", total=False)
    s.line("oc_enc", "Obligations — encours",
           lambda i, c: f"IF(AND({y(c)}>={I['annee_acq']},{y(c)}<={I['oc_ech']}),{I['oc']},0)", total=False)
    s.line("oc_int", "  Coupons obligataires", lambda i, c: f"{s.R('oc_enc', c)}*{I['oc_taux']}")
    s.line("oc_rmb", "  Remboursement obligations", lambda i, c: f"IF({y(c)}={I['oc_ech']},{I['oc']},0)")

    # Résultat net
    s.row += 1
    s.line("fin", "Charges financières", lambda i, c: f"-({s.R('d_int', c)}+{s.R('cv_int', c)}+{s.R('oc_int', c)})")
    s.line("rai", "Résultat avant impôt", lambda i, c: f"{s.R('ebit', c)}+{s.R('fin', c)}", bold=True)
    s.line("def", "Stock de déficits reportables (fin d'année)",
           lambda i, c: (f"MAX(0,-{s.R('rai', c)})" if i == 1 else
                         f"MAX(0,{s.R('def', s.prev(c))}-{s.R('rai', c)})"), total=False)
    s.line("is", "Impôt sur les sociétés",
           lambda i, c: (f"-MAX(0,{s.R('rai', c)})*{I['is_taux']}" if i == 1 else
                         f"-MAX(0,{s.R('rai', c)}-{s.R('def', s.prev(c))})*{I['is_taux']}"))
    s.line("rn", "Résultat net", lambda i, c: f"{s.R('rai', c)}+{s.R('is', c)}", bold=True)

    # BFR et flux
    s.row += 1
    s.put(s.row, 1, "Tableau des flux de trésorerie", font=BOLD)
    s.row += 1
    s.line("bfr", "BFR (niveau)", lambda i, c: f"{s.R('ca', c)}*{I['bfr_pct']}", total=False)
    s.line("dbfr", "Variation du BFR",
           lambda i, c: (f"-{s.R('bfr', c)}" if i == 1 else f"-({s.R('bfr', c)}-{s.R('bfr', s.prev(c))})"))
    s.line("cfo", "Flux d'exploitation (EBITDA − IS − ΔBFR)",
           lambda i, c: f"{s.R('ebitda', c)}+{s.R('is', c)}+{s.R('dbfr', c)}", bold=True)
    s.line("cfi", "Flux d'investissement",
           lambda i, c: f"-({s.R('acq', c)}+{s.R('capex_init', c)}+{s.R('ge', c)})", bold=True)
    s.line("apports", "  Apports en fonds propres + subventions",
           lambda i, c: f"IF({y(c)}=1,{I['apport_fondateur']},0)+IF({y(c)}={I['annee_acq']},{I['fp_invest']}*(1-{I['retrait_inv']})+{I['subventions']},0)")
    s.line("emprunts", "  Tirages dette + crédit-vendeur + obligations",
           lambda i, c: f"IF({y(c)}={I['annee_acq']},{I['dette']}+{I['cv']}+{I['oc']},0)")
    s.line("serv", "  Service de la dette (intérêts + capital + coupons + rembt obligations)",
           lambda i, c: f"-({s.R('d_int', c)}+{s.R('d_amo', c)}+{s.R('cv_int', c)}+{s.R('cv_amo', c)}+{s.R('oc_int', c)}+{s.R('oc_rmb', c)})")
    s.line("cff", "Flux de financement", lambda i, c: f"{s.R('apports', c)}+{s.R('emprunts', c)}+{s.R('serv', c)}", bold=True)
    s.line("var_treso", "Variation de trésorerie", lambda i, c: f"{s.R('cfo', c)}+{s.R('cfi', c)}+{s.R('cff', c)}", bold=True)
    s.line("treso", "Trésorerie fin d'année",
           lambda i, c: (s.R('var_treso', c) if i == 1 else f"{s.R('treso', s.prev(c))}+{s.R('var_treso', c)}"),
           total=False, bold=True)
    s.line("fcf_exp", "Cash-flow libre d'exploitation (CFO − gros entretien − service dette hors obligations in fine)",
           lambda i, c: f"{s.R('cfo', c)}-{s.R('ge', c)}-({s.R('d_int', c)}+{s.R('d_amo', c)}+{s.R('cv_int', c)}+{s.R('cv_amo', c)}+{s.R('oc_int', c)})")
    s.line("dscr", "DSCR ((EBITDA − IS) / service de la dette)",
           lambda i, c: f"IF(({s.R('d_int', c)}+{s.R('d_amo', c)}+{s.R('cv_int', c)}+{s.R('cv_amo', c)}+{s.R('oc_int', c)})=0,\"\","
                        f"({s.R('ebitda', c)}+{s.R('is', c)})/({s.R('d_int', c)}+{s.R('d_amo', c)}+{s.R('cv_int', c)}+{s.R('cv_amo', c)}+{s.R('oc_int', c)}))",
           fmt='0.00', total=False)

    # Investisseurs
    s.row += 1
    s.put(s.row, 1, "Rentabilité des investisseurs et contrôle", font=BOLD)
    s.row += 1
    s.line("dette_tot", "Dettes financières totales fin d'année",
           lambda i, c: f"{s.R('d_fin', c)}+{s.R('cv_fin', c)}+{s.R('oc_enc', c)}-{s.R('oc_rmb', c)}", total=False)
    s.line("v_actif", "Valeur estimée de l'actif (si cédé cette année)",
           lambda i, c: f"{s.R('possede', c)}*({I['prix']}+{s.R('cum_capex', c)})*{I['valo_sortie_pct']}", total=False)
    s.line("v_eq", "Valeur des fonds propres (actif + trésorerie − dettes)",
           lambda i, c: f"{s.R('v_actif', c)}+{s.R('treso', c)}-{s.R('dette_tot', c)}", total=False)
    s.line("flux_inv", "Flux des investisseurs (entrée an d'acq., sortie an 10)",
           lambda i, c: f"-IF({y(c)}={I['annee_acq']},{I['fp_invest']}*(1-{I['retrait_inv']}),0)"
                        + (f"+MAX(0,{s.R('v_eq', c)})*(1-{I['part_fondateur']})" if i == N else ""))
    s.line("flux_fond", "Flux du fondateur (apport an 1, sortie théorique an 10)",
           lambda i, c: f"-IF({y(c)}=1,{I['apport_fondateur']},0)"
                        + (f"+MAX(0,{s.R('v_eq', c)})*{I['part_fondateur']}" if i == N else ""))

    # Indicateurs clés
    s.section("8. Indicateurs clés")
    kpis = [
        ("k_besoin", "Emplois totaux (acq. + frais + travaux initiaux)",
         f"={I['prix']}*(1+{I['frais_acq_pct']})+SUM(C{I['capex_init']}:L{I['capex_init']})", EUR),
        ("k_ress", "Ressources externes et apports prévus",
         f"={I['apport_fondateur']}+{I['fp_invest']}*(1-{I['retrait_inv']})+{I['subventions']}+{I['dette']}+{I['cv']}+{I['oc']}", EUR),
        ("k_ca5", "CA année 5", f"=G{I['ca']}", EUR),
        ("k_ebitda5", "EBITDA année 5", f"=G{I['ebitda']}", EUR),
        ("k_ebitda10", "EBITDA année 10", f"=L{I['ebitda']}", EUR),
        ("k_equ5", "Part équestre du CA année 5", f"=G{I['part_equ']}", PCT),
        ("k_tmin", "Trésorerie minimale sur 10 ans", f"=MIN(C{I['treso']}:L{I['treso']})", EUR),
        ("k_t10", "Trésorerie fin année 10", f"=L{I['treso']}", EUR),
        ("k_dscr", "DSCR minimal (années 4 à 10)", f"=IF(COUNT(F{I['dscr']}:L{I['dscr']})=0,\"n/a\",MIN(F{I['dscr']}:L{I['dscr']}))", '0.00'),
        ("k_fcfneg", "Nb d'années à cash-flow libre d'exploitation négatif (an 4–10)",
         f"=COUNTIF(F{I['fcf_exp']}:L{I['fcf_exp']},\"<0\")", '0'),
        ("k_tri_inv", "TRI investisseurs (sortie théorique an 10)",
         f"=IFERROR(IRR(C{I['flux_inv']}:L{I['flux_inv']}),\"non calculable\")", PCT),
        ("k_tri_fond", "TRI fondateur (sortie théorique an 10)",
         f"=IFERROR(IRR(C{I['flux_fond']}:L{I['flux_fond']}),\"non calculable\")", PCT),
        ("k_part", "Part du capital du fondateur", f"={I['part_fondateur']}", PCT),
        ("k_votes", "Droits de vote du fondateur", f"={I['votes_fondateur']}", PCT),
    ]
    for key, lab, f, fmt in kpis:
        s.put(s.row, 1, lab)
        s.put(s.row, 2, f, fmt=fmt, font=BOLD)
        I[key] = f"$B${s.row}"
        s.row += 1
    s.put(s.row, 1, "Verdict (règle : trésorerie ≥ réserve cible, DSCR ≥ 1,2, cash-flow libre ≥ 0 an 4–10)", font=BOLD)
    s.put(s.row, 2,
          f"=IF(AND({I['k_tmin']}>=0,{I['k_t10']}>={I['treso_min']},OR({I['k_dscr']}=\"n/a\",{I['k_dscr']}>=1.2),{I['k_fcfneg']}=0),"
          f"\"VIABLE sous hypothèses\",IF({I['k_tmin']}<0,\"NON VIABLE (trésorerie négative)\",\"FRAGILE\"))", font=BOLD)
    I["k_verdict"] = f"$B${s.row}"
    s.row += 1
    ws.freeze_panes = "B2"
    return s


def build():
    wb = Workbook()
    readme = wb.active
    readme.title = "Lisez-moi"
    readme.column_dimensions["A"].width = 120
    for i, line in enumerate(H.LISEZMOI, start=1):
        c = readme.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True)
        if i == 1:
            c.font = Font(bold=True, size=14)

    sheets = []
    for sheet_name, params in H.feuilles():
        sheets.append((sheet_name, params, build_scenario(wb, sheet_name, params)))

    comp = wb.create_sheet("Comparaison", 1)
    comp.column_dimensions["A"].width = 30
    keys = [("k_besoin", "Emplois totaux", EUR), ("k_ress", "Ressources prévues", EUR),
            ("k_ca5", "CA an 5", EUR), ("k_ebitda5", "EBITDA an 5", EUR), ("k_ebitda10", "EBITDA an 10", EUR),
            ("k_equ5", "Part équestre CA an 5", PCT), ("k_tmin", "Trésorerie min.", EUR),
            ("k_t10", "Trésorerie an 10", EUR), ("k_dscr", "DSCR min. an 4–10", '0.00'),
            ("k_fcfneg", "Années FCF < 0", '0'), ("k_tri_inv", "TRI investisseurs", PCT),
            ("k_tri_fond", "TRI fondateur", PCT), ("k_part", "Capital fondateur", PCT),
            ("k_votes", "Votes fondateur", PCT), ("k_verdict", "Verdict", None)]
    comp.cell(row=1, column=1, value="Comparaison des scénarios et stress tests (valeurs liées aux feuilles)").font = Font(bold=True, size=13)
    comp.cell(row=3, column=1, value="Feuille").font = BOLD
    for j, (_, lab, _) in enumerate(keys):
        comp.cell(row=3, column=2 + j, value=lab).font = BOLD
        comp.column_dimensions[get_column_letter(2 + j)].width = 16
    for i, (name, params, s) in enumerate(sheets):
        comp.cell(row=4 + i, column=1, value=name)
        for j, (k, _, fmt) in enumerate(keys):
            c = comp.cell(row=4 + i, column=2 + j, value=f"='{name}'!{s.ref[k]}")
            if fmt:
                c.number_format = fmt
    comp.freeze_panes = "B4"

    src = wb.create_sheet("Sources_hypotheses")
    src.column_dimensions["A"].width = 60
    src.column_dimensions["B"].width = 100
    src.cell(row=1, column=1, value="Hypothèse").font = BOLD
    src.cell(row=1, column=2, value="Source / justification / statut").font = BOLD
    for i, (a, b) in enumerate(H.SOURCES, start=2):
        src.cell(row=i, column=1, value=a)
        src.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True)

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Previsionnels_10_ans.xlsx")
    wb.save(out)
    print("écrit :", out, "—", len(wb.sheetnames), "feuilles")


if __name__ == "__main__":
    build()
