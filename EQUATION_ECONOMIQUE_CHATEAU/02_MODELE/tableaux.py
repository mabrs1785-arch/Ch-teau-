"""Génère les tableaux Markdown des livrables à partir du classeur recalculé.

Les nombres des livrables viennent donc tous du modèle : aucune recopie à la main.
Usage : python3 tableaux.py  →  écrit 02_MODELE/sorties/tableaux.md
"""
from __future__ import annotations

import json
import os

from openpyxl import load_workbook

import moteurs  # noqa: F401  (registre des hypothèses, pour les investissements saisis comme hypothèses)
import parametres  # noqa: F401

ICI = os.path.dirname(os.path.abspath(__file__))
XL = os.path.join(ICI, "sorties", "Modele_equation_chateau.xlsx")


def k(v):
    """Montant en k€ arrondi, avec séparateur de milliers français."""
    return f"{v / 1000:,.0f}".replace(",", " ") + " k€"


def m(v, d=2):
    return f"{v / 1e6:,.{d}f}".replace(",", " ").replace(".", ",") + " M€"


def pct(v):
    return f"{v * 100:.1f}".replace(".", ",") + " %"


def lire(ws, col_debut=5, nb=9):
    d = {}
    for r in range(5, ws.max_row + 1):
        lid = ws.cell(row=r, column=1).value
        if lid:
            d[lid] = [ws.cell(row=r, column=c).value for c in range(col_debut, col_debut + nb)]
    return d


def tableau_budget(b):
    cats = [
        ("Bâti : entretien courant", ["bati_courant"]),
        ("Bâti : provision de gros entretien, aléas de chantier et maîtrise d'œuvre", ["bati_ge", "imprevus_travaux", "moe_do"]),
        ("Sécurité", ["securite"]),
        ("Personnel", ["dont_personnel_calc"]),
        ("Jardins, parc, prairies, clôtures, bois (hors personnel)", ["jardins_reguliers", "parc", "prairies", "clotures", "bois"]),
        ("Équipements (renouvellement, entretien, carburant)", ["equip_renouv", "equip_entretien"]),
        ("Chevaux du propriétaire et installations équestres", ["chevaux", "installations_eq"]),
        ("Énergie et eau", ["chauffage", "hors_gel", "electricite", "eau"]),
        ("Assurances", ["assurance_bati", "assurance_autres"]),
        ("Impôts et taxes (foncières, IFI)", ["taxe_fonciere_bati", "taxe_fonciere_non_bati", "ifi"]),
        ("Administration et suivi patrimonial", ["admin", "honoraires_patrimoine"]),
        ("Imprévus d'exploitation", ["imprevus"]),
        ("Investissements futurs", ["invest_futurs"]),
    ]
    lignes = ["| Poste (€ 2026 TTC par an, cas central) | S1 | S2 | S3 |", "|---|---:|---:|---:|"]
    for lib, ids in cats:
        vals = [sum(b[i][j] for i in ids) for j in (1, 4, 7)]
        lignes.append(f"| {lib} | {k(vals[0])} | {k(vals[1])} | {k(vals[2])} |")
    t = b["total_besoin"]
    lignes.append(f"| **Besoin annuel total (central)** | **{k(t[1])}** | **{k(t[4])}** | **{k(t[7])}** |")
    lignes.append(f"| Fourchette favorable – défavorable | {k(t[0])} – {k(t[2])} | {k(t[3])} – {k(t[5])} | {k(t[6])} – {k(t[8])} |")
    lignes.append(f"| Effectif salarié (ETP) | {b['etp_total'][1]:.1f} | {b['etp_total'][4]:.1f} | {b['etp_total'][7]:.1f} |".replace(".", ","))
    lignes.append(f"| Part du personnel / part du bâti | {pct(b['part_personnel'][1])} / {pct(b['part_bati'][1])} | "
                  f"{pct(b['part_personnel'][4])} / {pct(b['part_bati'][4])} | {pct(b['part_personnel'][7])} / {pct(b['part_bati'][7])} |")
    lignes.append(f"| Valeur de reconstruction retenue | {m(b['valeur_reconstruction'][1], 1)} | {m(b['valeur_reconstruction'][4], 1)} | {m(b['valeur_reconstruction'][7], 1)} |")
    lignes.append(f"| Besoin en % de la valeur de reconstruction | {pct(b['besoin_en_pct_valeur'][1])} | {pct(b['besoin_en_pct_valeur'][4])} | {pct(b['besoin_en_pct_valeur'][7])} |")
    lignes.append(f"| Besoin par m² bâti | {b['besoin_par_m2'][1]:.0f} €/m² | {b['besoin_par_m2'][4]:.0f} €/m² | {b['besoin_par_m2'][7]:.0f} €/m² |")
    lignes.append(f"| Si porté par une société assujettie (TVA récupérée) | {k(b['besoin_ht_societe'][1])} | {k(b['besoin_ht_societe'][4])} | {k(b['besoin_ht_societe'][7])} |")
    lignes.append(f"| Scénario complémentaire : + annuité de dette (70 % du prix, 20 ans, 4,2 %) | {k(b['dette_annuite'][1])} | {k(b['dette_annuite'][4])} | {k(b['dette_annuite'][7])} |")
    lignes.append(f"| Total avec dette | {k(b['total_avec_dette'][1])} | {k(b['total_avec_dette'][4])} | {k(b['total_avec_dette'][7])} |")
    return "\n".join(lignes)


def tableau_moteurs(mo):
    noms = [("eq_std_cn", "Pension standard (30 boxes)", "B"), ("eq_hdg_cn", "Pension haut de gamme (30 boxes)", "B"),
            ("eq_pro_cn", "Écuries louées à un professionnel (20 boxes)", "A"), ("eq_ret_cn", "Pension retraite au pré (30 chevaux)", "B"),
            ("repro_cn", "Centre de reproduction (si vétérinaire et étalon)", "B"), ("evt_cn", "Réceptions privées plafonnées (10/an)", "C"),
            ("tourn_cn", "Tournages et prises de vue (2 j/an)", "C"), ("semj_cn", "Séminaires de jour (20 j/an)", "C"),
            ("gite_cn", "Grande maison louée à la semaine", "C"), ("log_cn", "Logements locatifs dans les communs (600 m²)", "A"),
            ("stock_cn", "Hivernage de camping-cars (600 m²)", "A"), ("auto_cn", "Gardiennage de voitures de collection", "A"),
            ("pv_cn", "Bail solaire (5 ha hors champ de visibilité)", "A"), ("bois_cn", "Forêt (30 ha)", "A"),
            ("hot_cn", "Château-hôtel opéré (15 chambres)", "B/C"),
            ("anc_cn", "Bail d'ancrage à un exploitant de séminaires (50 chambres, si conditions réunies)", "A"),
            ("lux_cn", "Location de luxe de la partie privée (3 semaines par an)", "C")]
    inv = {"evt_cn": "evt_inv", "gite_cn": "gite_inv", "log_cn": "log_inv", "hot_cn": "hot_inv",
           "stock_cn": "stock_capex", "auto_cn": "auto_capex", "lux_cn": "lux_capex"}
    lignes = ["| Moteur (échelle S2) | Classe | Favorable | Central | Défavorable | Investissement central |",
              "|---|:---:|---:|---:|---:|---:|"]
    for lid, lib, cl in noms:
        v = mo[lid]
        if lid not in inv:
            i = "—"
        elif inv[lid] in mo:
            i = k(mo[inv[lid]][1])
        else:   # investissement saisi comme hypothèse (feuille Hypotheses), pas comme ligne calculée
            from noyau import REGISTRE
            i = k(REGISTRE[inv[lid]].cen) if inv[lid] in REGISTRE else "—"
        lignes.append(f"| {lib} | {cl} | {k(v[0])} | {k(v[1])} | {k(v[2])} | {i} |")
    return "\n".join(lignes)


def tableau_modeles(md):
    titres = {"M1": "Socle contractuel et pension retraite", "M2": "Lieu d'exception à quota",
              "M3": "Équestre haut de gamme et services", "M4": "Château-hôtel opéré + réceptions",
              "M5": "Socle minimal + dotation", "M6": "Socle minimal + entreprise extérieure",
              "M7": "Locataire d'ancrage + socle foncier", "M8": "Socle M1 + location de luxe"}
    lignes = ["| Modèle | Contribution disponible (central) | Couverture du besoin S2 (fav / cen / déf) | Écart à financer (central) | Dotation équivalente (central) | EBITDA équivalent (central) | Investissement |",
              "|---|---:|---:|---:|---:|---:|---:|"]
    for mid, t in titres.items():
        c = md[f"couv_{mid}"]
        lignes.append(f"| {mid} — {t} | {k(md[f'{mid}_nette'][1])} | {pct(c[0])} / {pct(c[1])} / {pct(c[2])} | "
                      f"{k(md[f'{mid}_ecart'][1])} | {m(md[f'{mid}_dotation'][1], 1)} | {m(md[f'{mid}_ebitda'][1], 1)} | {k(md[f'{mid}_invest'][1])} |")
    return "\n".join(lignes)


def tableau_stress(ws):
    d = {}
    for r in range(5, ws.max_row + 1):
        lid = ws.cell(row=r, column=1).value
        if lid:
            d[lid] = ws.cell(row=r, column=5).value
    feuilles = [x[: -len("_min_res")] for x in d if x.endswith("_min_res")]
    lignes = ["| Projection (20 ans) | Réserve minimale | Réserve en année 20 | Dotation en année 20 | Années en déficit | Années en tension | Contribution moyenne des activités |",
              "|---|---:|---:|---:|---:|---:|---:|"]
    for f in feuilles:
        lignes.append(f"| {f} | {m(d[f + '_min_res'])} | {m(d[f + '_res20'])} | {m(d[f + '_dot20'], 1)} | {d[f + '_def']:.0f} | "
                      f"{d.get(f + '_tension', 0):.0f} | {k(d[f + '_act'])} |")
    return "\n".join(lignes)


def main():
    wb = load_workbook(XL, data_only=True)
    b = lire(wb["Budget"])
    mo = lire(wb["Moteurs"], nb=3)
    md = lire(wb["Modeles"], nb=3)
    out = ["# Tableaux générés depuis le modèle (ne pas modifier à la main)", "",
           "## Budget annuel", "", tableau_budget(b), "",
           "## Moteurs de revenus", "", tableau_moteurs(mo), "",
           "## Modèles", "", tableau_modeles(md), "",
           "## Projections et stress tests", "", tableau_stress(wb["Stress_synthese"]), ""]
    for f in ("mc_dimensionnement.json", "mc_dotation_necessaire.json", "mc_grille_hybride.json",
              "mc_holding_hybride.json", "mc_ruptures_S2.json", "mc_ruptures_S1.json", "mc_ruptures_S3.json",
              "mc_distribution_solutions.json", "mc_comparaison_modeles.json", "mc_sensibilite_rendement.json"):
        chemin = os.path.join(ICI, "sorties", f)
        if os.path.exists(chemin):
            out += [f"## {f}", "", "```json", open(chemin, encoding="utf-8").read(), "```", ""]
    with open(os.path.join(ICI, "sorties", "tableaux.md"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(out))
    print("\n".join(out[:60]))


if __name__ == "__main__":
    main()
