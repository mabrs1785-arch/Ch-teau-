"""Construit le classeur Excel du modèle, le recalcule avec LibreOffice et vérifie chaque formule.

Usage : python3 construire_modele.py
Sortie : 02_MODELE/sorties/Modele_equation_chateau.xlsx et verification.json
"""
from __future__ import annotations

import json
import os
import subprocess
import sys

from openpyxl import Workbook, load_workbook

ICI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ICI)

import parametres  # noqa: E402,F401  (remplit le registre)
from budget import LIGNES as LIGNES_BUDGET  # noqa: E402
from excel_outils import (F_NORMAL, F_TITRE, ecrire_hypotheses, ecrire_sources,  # noqa: E402
                          ecrire_tableau_gabarits)
from noyau import LIBELLES_CAS, REGISTRE, noms_utilises  # noqa: E402

RACINE = os.path.dirname(ICI)
SORTIES = os.path.join(ICI, "sorties")
FICHIER = os.path.join(SORTIES, "Modele_equation_chateau.xlsx")
RECALC = ("/root/.claude/skills/synced/031d3383-3680-4a27-8446-9041f5912897_"
          "70cae721-469c-4e22-9a68-fb55dde9cb70/xlsx/scripts/recalc.py")

MODULES_OPTIONNELS = ["moteurs", "modeles", "projection"]


def verifier_noms(lignes, scenarios):
    manquants = set()
    for (_lid, _lib, gab, *_r) in lignes:
        for nom in noms_utilises(gab):
            for s in scenarios:
                n = nom.replace("{S}", s) if "{S}" in nom else nom
                for suffixe in ("", "_fav", "_def"):
                    n2 = n.replace("{C}", suffixe)
                    if n2 not in REGISTRE and not any(
                            n2 == pid + suf for pid in REGISTRE for suf in ("_fav", "_def")):
                        manquants.add(n2)
    return sorted(manquants)


def lisez_moi(wb):
    ws = wb.active
    ws.title = "Lisez-moi"
    lignes = [
        ("MODÈLE — L'ÉQUATION ÉCONOMIQUE D'UN CHÂTEAU", F_TITRE),
        ("", F_NORMAL),
        ("Objet : budget annuel de fonctionnement de trois domaines types (S1, S2, S3), "
         "contribution nette des moteurs de revenus, comparaison des modèles économiques, "
         "projection sur 20 ans de la meilleure solution, stress tests.", F_NORMAL),
        ("Unités : € constants 2026, TTC pour un propriétaire particulier, sauf mention contraire.", F_NORMAL),
        ("Couleurs : bleu = saisie modifiable (feuille Hypotheses) ; noir = formule ; "
         "fond jaune = hypothèse clé ; vert = lien vers une autre feuille.", F_NORMAL),
        ("Statuts : SOURCÉ (source citée, voir feuille Sources), ESTIMATION (raisonnement sur sources), "
         "HYPOTHÈSE (choix de modélisation à valider).", F_NORMAL),
        ("Trois cas par hypothèse : favorable, central, défavorable pour l'équilibre du domaine.", F_NORMAL),
        ("Toutes les cellules calculées sont des formules : modifier une hypothèse recalcule tout.", F_NORMAL),
        ("Méthode, définitions et limites : voir 00_PILOTAGE/Cadre_analytique.md et 03_LIVRABLES/.", F_NORMAL),
        ("Vérification : chaque formule a été recalculée par LibreOffice et comparée à un calcul Python "
         "indépendant (02_MODELE/sorties/verification.json).", F_NORMAL),
    ]
    for i, (txt, police) in enumerate(lignes, start=1):
        c = ws.cell(row=i, column=1, value=txt)
        c.font = police
    ws.column_dimensions["A"].width = 150


def construire():
    os.makedirs(SORTIES, exist_ok=True)
    scenarios = ["S1", "S2", "S3"]
    manquants = verifier_noms(LIGNES_BUDGET, scenarios)
    if manquants:
        raise SystemExit(f"Hypothèses manquantes dans le registre : {manquants}")

    modules = []
    for nom_module in MODULES_OPTIONNELS:
        try:
            modules.append(__import__(nom_module))
        except ModuleNotFoundError:
            continue
    wb = Workbook()
    lisez_moi(wb)
    ecrire_hypotheses(wb, ids_cles=getattr(parametres, "IDS_CLES", ()))
    colonnes = [(f"{s} {LIBELLES_CAS[c]}", s, c) for s in scenarios for c in ("fav", "cen", "def")]
    _, controle = ecrire_tableau_gabarits(
        wb, "Budget", "BUDGET ANNUEL DE FONCTIONNEMENT — trois domaines types, trois cas (€ 2026 TTC)",
        LIGNES_BUDGET, colonnes, prefixe_noms="bud_",
        note=("Chaque cellule est une formule construite à partir des hypothèses nommées "
              "(feuille Hypotheses). La provision de gros entretien applique un taux à la valeur de "
              "reconstruction (règle du NRC) ; la méthode par composants n'a pas pu être sourcée. Les honoraires "
              "de maîtrise d'œuvre et l'assurance dommages-ouvrage sont ajoutés en pourcentage du gros entretien."))

    for module in modules:
        controle.update(module.ecrire(wb))

    ecrire_sources(wb, os.path.join(RACINE, "01_RECHERCHE", "sources"))
    wb.save(FICHIER)
    return controle


def recalculer_et_verifier(controle, tolerance=0.5):
    res = subprocess.run([sys.executable, RECALC, FICHIER, "120"], capture_output=True, text=True)
    rapport = json.loads(res.stdout)
    wb = load_workbook(FICHIER, data_only=True)
    ecarts = []
    for (feuille, cellule), attendu in controle.items():
        obtenu = wb[feuille][cellule].value
        if obtenu is None or abs((obtenu or 0) - attendu) > tolerance + 1e-9 * abs(attendu):
            ecarts.append({"feuille": feuille, "cellule": cellule, "python": attendu, "excel": obtenu})
    import monte_carlo   # contrôle du jumeau : la simulation sans chocs doit reproduire les projections
    sortie = {"recalcul": rapport, "cellules_comparees": len(controle),
              "ecarts": ecarts[:50], "nb_ecarts": len(ecarts),
              "jumeau_monte_carlo_ecart_max_eur": monte_carlo.verifier_jumeau(FICHIER)}
    with open(os.path.join(SORTIES, "verification.json"), "w", encoding="utf-8") as f:
        json.dump(sortie, f, ensure_ascii=False, indent=2)
    return sortie


if __name__ == "__main__":
    ctrl = construire()
    s = recalculer_et_verifier(ctrl)
    print(json.dumps({k: s[k] for k in ("cellules_comparees", "nb_ecarts")}, ensure_ascii=False),
          json.dumps(s["recalcul"], ensure_ascii=False)[:300])
    if s["nb_ecarts"]:
        print("ÉCARTS :", json.dumps(s["ecarts"][:10], ensure_ascii=False, indent=1))
    print("Jumeau Monte-Carlo sans chocs, écart maximal (€) :",
          json.dumps(s["jumeau_monte_carlo_ecart_max_eur"], ensure_ascii=False))
